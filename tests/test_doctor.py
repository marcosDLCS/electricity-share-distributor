"""Unit tests for the DataDoctor diagnostic engine and CLI doctor command."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
from typer.testing import CliRunner

from src.cli import app
from src.ingestion.doctor import DataDoctor

runner = CliRunner()


def test_doctor_missing_directories(tmp_path: Path) -> None:
    non_existent = tmp_path / "does_not_exist"
    doc = DataDoctor(consumption_dir=non_existent, generation_dir=non_existent)
    report = doc.diagnose()
    assert report.overall_status == "error"
    assert not report.can_calculate


def test_doctor_detects_unexpected_gaps(tmp_path: Path) -> None:
    c_dir = tmp_path / "consumption"
    g_dir = tmp_path / "generation"
    c_dir.mkdir()
    g_dir.mkdir()

    # Create 48-hour consumption CSV with 5 hours missing
    # Hours 01:00 to 24:00 on Day 1, but Day 2 skips hours 10:00 to 14:00
    dates = pd.date_range("2026-05-01 00:00:00", "2026-05-02 23:00:00", freq="h")
    rows = []
    for dt in dates:
        if dt.day == 2 and 10 <= dt.hour <= 14:
            continue  # Skip 5 hours
        datadis_hour = f"{dt.hour + 1:02d}:00"
        date_str = dt.strftime("%Y/%m/%d")
        rows.append(f"ES0021000000000001AA;{date_str};{datadis_hour};0,250;IBERDROLA;R")

    (c_dir / "ES0021000000000001AA.csv").write_text(
        "CUPS;Fecha;Hora;Consumo_kWh;Metodo;Tipo\n" + "\n".join(rows),
        encoding="utf-8",
    )

    # Huawei Excel covering the same 48 hours without gaps
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Informe de plantas", "", ""])
    ws.append(["Período estadístico", "Rendimiento FV (kWh)", "Rendimiento del inversor (kWh)"])
    for dt in dates:
        ws.append([dt.strftime("%Y-%m-%d %H:%M:%S"), 2.5, 2.5])
    wb.save(g_dir / "generation.xlsx")

    doc = DataDoctor(consumption_dir=c_dir, generation_dir=g_dir)
    report = doc.diagnose()

    assert report.overall_status == "warning"
    assert len(report.consumption_results) == 1
    c_res = report.consumption_results[0]
    assert c_res.missing_hours == 5
    assert len(c_res.gaps) == 1
    assert c_res.gaps[0].missing_hours == 5


def test_doctor_detects_inactive_meter(tmp_path: Path) -> None:
    c_dir = tmp_path / "consumption"
    g_dir = tmp_path / "generation"
    c_dir.mkdir()
    g_dir.mkdir()

    dates = pd.date_range("2026-05-01 00:00:00", "2026-05-01 23:00:00", freq="h")
    rows = []
    for dt in dates:
        datadis_hour = f"{dt.hour + 1:02d}:00"
        date_str = dt.strftime("%Y/%m/%d")
        rows.append(f"ES0021000000000002BB;{date_str};{datadis_hour};0,000;IBERDROLA;R")

    (c_dir / "inactive.csv").write_text(
        "CUPS;Fecha;Hora;Consumo_kWh;Metodo;Tipo\n" + "\n".join(rows),
        encoding="utf-8",
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Informe de plantas", "", ""])
    ws.append(["Período estadístico", "Rendimiento FV (kWh)", "Rendimiento del inversor (kWh)"])
    for dt in dates:
        ws.append([dt.strftime("%Y-%m-%d %H:%M:%S"), 1.0, 1.0])
    wb.save(g_dir / "gen.xlsx")

    doc = DataDoctor(consumption_dir=c_dir, generation_dir=g_dir)
    report = doc.diagnose()

    assert report.overall_status == "warning"
    c_res = report.consumption_results[0]
    assert c_res.zero_ratio_pct == 100.0
    assert any("Inactive meter" in note for note in c_res.notes)


def test_cli_doctor_command() -> None:
    res = runner.invoke(app, ["doctor", "--lang", "en"])
    assert res.exit_code == 0
    assert "ESD DATA DOCTOR DIAGNOSTIC REPORT" in res.output
    assert "Participating CUPS" in res.output
    assert "Consumption Data Health" in res.output
    assert "Open-Meteo Hourly Weather" in res.output
    # Verify color tags are parsed and not output as literal markup tags
    assert "[green]" not in res.output
    assert "[/green]" not in res.output
    assert "[yellow]" not in res.output
    assert "[/yellow]" not in res.output


def test_cli_doctor_verbose() -> None:
    res = runner.invoke(app, ["doctor", "--verbose", "--lang", "en"])
    assert res.exit_code == 0
    assert "ESD DATA DOCTOR DIAGNOSTIC REPORT" in res.output


def test_doctor_cups_use_enrichment(tmp_path: Path) -> None:
    """Ensure DataDoctor loads CUPS configuration and attaches uses to SeriesCheckResult."""
    c_dir = tmp_path / "consumption"
    g_dir = tmp_path / "generation"
    c_dir.mkdir()
    g_dir.mkdir()

    dates = pd.date_range("2026-05-01 00:00:00", "2026-05-01 23:00:00", freq="h")
    rows = []
    for dt in dates:
        datadis_hour = f"{dt.hour + 1:02d}:00"
        date_str = dt.strftime("%Y/%m/%d")
        rows.append(f"ES0021000000000001AA;{date_str};{datadis_hour};1,250;IBERDROLA;R")

    (c_dir / "cups1.csv").write_text(
        "CUPS;Fecha;Hora;Consumo_kWh;Metodo;Tipo\n" + "\n".join(rows),
        encoding="utf-8",
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Informe de plantas", "", ""])
    ws.append(["Período estadístico", "Rendimiento FV (kWh)", "Rendimiento del inversor (kWh)"])
    for dt in dates:
        ws.append([dt.strftime("%Y-%m-%d %H:%M:%S"), 2.0, 2.0])
    wb.save(g_dir / "gen.xlsx")

    # Create synthetic cups.json
    cups_file = tmp_path / "cups.json"
    cups_file.write_text(
        '[{"cups": "ES0021000000000001AA", "use": "Central Heat Pump", "primary": true}]',
        encoding="utf-8",
    )

    doc = DataDoctor(
        consumption_dir=c_dir,
        generation_dir=g_dir,
        cups_config_path=cups_file,
    )
    report = doc.diagnose()

    assert len(report.consumption_results) == 1
    c_res = report.consumption_results[0]
    assert c_res.identifier == "ES0021000000000001AA"
    assert c_res.use == "Central Heat Pump"
    assert c_res.is_primary is True


def test_doctor_weather_checks(tmp_path: Path) -> None:
    """Ensure DataDoctor audits weather CSV files and checks time-series continuity."""
    c_dir = tmp_path / "consumption"
    g_dir = tmp_path / "generation"
    w_dir = tmp_path / "weather"
    c_dir.mkdir()
    g_dir.mkdir()
    w_dir.mkdir()

    dates = pd.date_range("2026-05-01 00:00:00", "2026-05-02 23:00:00", freq="h")
    c_rows = []
    for dt in dates:
        datadis_hour = f"{dt.hour + 1:02d}:00"
        date_str = dt.strftime("%Y/%m/%d")
        c_rows.append(f"ES0021000000000001AA;{date_str};{datadis_hour};0,500;IBERDROLA;R")
    (c_dir / "cups.csv").write_text(
        "CUPS;Fecha;Hora;Consumo_kWh;Metodo;Tipo\n" + "\n".join(c_rows),
        encoding="utf-8",
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Informe de plantas", "", ""])
    ws.append(["Período estadístico", "Rendimiento FV (kWh)", "Rendimiento del inversor (kWh)"])
    for dt in dates:
        ws.append([dt.strftime("%Y-%m-%d %H:%M:%S"), 3.0, 3.0])
    wb.save(g_dir / "gen.xlsx")

    # Clean weather CSV covering same 48 hours
    w_rows = ["timestamp,temperature,cloud_cover,weather_code"]
    for dt in dates:
        w_rows.append(f"{dt.strftime('%Y-%m-%d %H:%M:%S')},18.5,10,1")
    (w_dir / "2026_weather.csv").write_text("\n".join(w_rows), encoding="utf-8")

    doc = DataDoctor(
        consumption_dir=c_dir,
        generation_dir=g_dir,
        weather_dir=w_dir,
    )
    report = doc.diagnose()

    assert report.weather_result is not None
    assert report.weather_result.status == "ok"
    assert report.weather_result.total_records == 48
    assert report.weather_result.missing_hours == 0
    assert report.overall_status == "ok"


def test_doctor_weather_incomplete_coverage(tmp_path: Path) -> None:
    """Ensure DataDoctor flags warning when weather data does not cover optimization window."""
    c_dir = tmp_path / "consumption"
    g_dir = tmp_path / "generation"
    w_dir = tmp_path / "weather"
    c_dir.mkdir()
    g_dir.mkdir()
    w_dir.mkdir()

    dates = pd.date_range("2026-05-01 00:00:00", "2026-05-02 23:00:00", freq="h")
    c_rows = []
    for dt in dates:
        datadis_hour = f"{dt.hour + 1:02d}:00"
        date_str = dt.strftime("%Y/%m/%d")
        c_rows.append(f"ES0021000000000001AA;{date_str};{datadis_hour};0,500;IBERDROLA;R")
    (c_dir / "cups.csv").write_text(
        "CUPS;Fecha;Hora;Consumo_kWh;Metodo;Tipo\n" + "\n".join(c_rows),
        encoding="utf-8",
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Informe de plantas", "", ""])
    ws.append(["Período estadístico", "Rendimiento FV (kWh)", "Rendimiento del inversor (kWh)"])
    for dt in dates:
        ws.append([dt.strftime("%Y-%m-%d %H:%M:%S"), 3.0, 3.0])
    wb.save(g_dir / "gen.xlsx")

    # Incomplete weather CSV: only covers Day 1 (first 24h)
    w_rows = ["timestamp,temperature,cloud_cover,weather_code"]
    for dt in dates[:24]:
        w_rows.append(f"{dt.strftime('%Y-%m-%d %H:%M:%S')},18.5,10,1")
    (w_dir / "2026_weather.csv").write_text("\n".join(w_rows), encoding="utf-8")

    doc = DataDoctor(
        consumption_dir=c_dir,
        generation_dir=g_dir,
        weather_dir=w_dir,
    )
    report = doc.diagnose()

    assert report.weather_result is not None
    assert report.weather_result.status == "warning"
    assert any("does not fully cover" in n for n in report.weather_result.notes)
    assert report.overall_status == "warning"
    assert report.can_calculate is True


def test_cli_doctor_custom_flags(tmp_path: Path) -> None:
    """Ensure doctor CLI accepts --cups-file and --weather-dir options."""
    cups_file = tmp_path / "custom_cups.json"
    cups_file.write_text("[]", encoding="utf-8")
    w_dir = tmp_path / "custom_weather"
    w_dir.mkdir()

    res = runner.invoke(
        app, ["doctor", "--cups-file", str(cups_file), "--weather-dir", str(w_dir), "--lang", "en"]
    )
    assert res.exit_code == 0
    assert "ESD DATA DOCTOR DIAGNOSTIC REPORT" in res.output
