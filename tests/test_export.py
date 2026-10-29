"""Unit tests for report exports (monthly shares CSV, hourly shares CSV, and interactive HTML)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from src.ingestion.schema import IngestionSummary
from src.optimization.models import (
    CommunityMonthlyMetrics,
    CupsMonthlyMetrics,
    HourlyOptimizationResult,
    HourlyProfileMatrix,
    OptimizationResult,
    OptimizationStrategy,
)
from src.presentation.export import (
    FORMULA_INJECTION_PREFIXES,
    ReportExporter,
    sanitize_formula_injection,
)
from src.presentation.views.common import MONTH_ABBR_EN, MONTH_ABBR_ES, format_month_abbr


def create_dummy_optimization_result() -> tuple[OptimizationResult, IngestionSummary]:
    cups_list = ["ES0021000000000001AA", "ES0021000000000002BB"]
    c1 = CupsMonthlyMetrics(
        cups="ES0021000000000001AA",
        use="Aerotermia",
        month="2026-05",
        beta=0.6,
        consumption_kwh=500.0,
        generation_allocated_kwh=600.0,
        self_consumed_kwh=480.0,
        surplus_kwh=120.0,
        grid_demand_kwh=20.0,
        self_consumption_rate=80.0,
        solar_coverage_rate=96.0,
    )
    c2 = CupsMonthlyMetrics(
        cups="ES0021000000000002BB",
        use="Garaje",
        month="2026-05",
        beta=0.4,
        consumption_kwh=300.0,
        generation_allocated_kwh=400.0,
        self_consumed_kwh=270.0,
        surplus_kwh=130.0,
        grid_demand_kwh=30.0,
        self_consumption_rate=67.5,
        solar_coverage_rate=90.0,
    )
    m1 = CommunityMonthlyMetrics(
        month="2026-05",
        strategy=OptimizationStrategy.OPTIMAL,
        total_generation_kwh=1000.0,
        total_consumption_kwh=800.0,
        total_self_consumed_kwh=750.0,
        total_surplus_kwh=250.0,
        total_grid_demand_kwh=50.0,
        self_consumption_rate=75.0,
        solar_coverage_rate=93.75,
        beta_sum=1.0,
        cups_metrics=[c1, c2],
    )

    hours = [f"{h:02d}" for h in range(24)]
    profile_matrix = HourlyProfileMatrix(
        cups_list=cups_list,
        hours=hours,
        matrix={
            cups_list[0]: dict.fromkeys(hours, 0.6),
            cups_list[1]: dict.fromkeys(hours, 0.4),
        },
        hourly_sums=dict.fromkeys(hours, 1.0),
        average_shares={cups_list[0]: 0.6, cups_list[1]: 0.4},
        average_sum=1.0,
    )
    ts_rows = []
    for h in hours:
        ts_rows.append(
            {
                "timestamp": f"2026-05-01 {h}:00:00",
                "hour": int(h),
                "cups": cups_list[0],
                "use": "Aerotermia",
                "consumption_kwh": 2.4,
                "generation_kwh": 5.0,
                "share_pct": "60%",
                "estimated_share_pct": "60%",
                "solar_efficiency": 75,
            }
        )
        ts_rows.append(
            {
                "timestamp": f"2026-05-01 {h}:00:00",
                "hour": int(h),
                "cups": cups_list[1],
                "use": "Garaje",
                "consumption_kwh": 1.6,
                "generation_kwh": 5.0,
                "share_pct": "40%",
                "estimated_share_pct": "40%",
                "solar_efficiency": 75,
            }
        )

    result = OptimizationResult(
        strategy=OptimizationStrategy.OPTIMAL,
        monthly_results=[m1],
        total_summary=m1,
        baselines={},
        hourly_result=HourlyOptimizationResult(
            profile_matrix=profile_matrix,
            time_series_df=pd.DataFrame(ts_rows),
        ),
        cups_uses={"ES0021000000000001AA": "Aerotermia", "ES0021000000000002BB": "Garaje"},
        current_self_consumed_kwh=480.0,
        monthly_share_self_consumed_kwh=750.0,
        hourly_share_self_consumed_kwh=820.0,
    )

    summary = IngestionSummary(
        cups_count=2,
        cups_list=cups_list,
        start_time="2026-05-01 00:00:00",
        end_time="2026-05-31 23:00:00",
        total_hours=744,
        consumption_files_loaded=2,
        generation_files_loaded=1,
    )

    return result, summary


def test_export_monthly_shares_csv(tmp_path: Path) -> None:
    opt_result, _ = create_dummy_optimization_result()
    csv_path = ReportExporter(
        output_dir=tmp_path, timestamp="20260531_120000"
    ).export_monthly_shares_csv(opt_result, precision=1)
    assert csv_path.exists()
    assert csv_path.name == "20260531_120000_esd_monthly_shares.csv"

    df = pd.read_csv(csv_path, sep=";", decimal=",")
    assert len(df) == 2
    expected_cols = [
        "month",
        "cups",
        "use",
        "consumption_kwh",
        "generation_kwh",
        "share_pct",
        "self_consumed_kwh",
        "surplus_kwh",
        "grid_demand_kwh",
    ]
    assert list(df.columns) == expected_cols
    assert df.iloc[0]["use"] == "Aerotermia"
    assert df.iloc[0]["share_pct"] == "60,0%"
    assert df.iloc[0]["consumption_kwh"] == 500.0
    assert df.iloc[0]["generation_kwh"] == 1000.0
    assert df.iloc[0]["self_consumed_kwh"] == 480.0
    assert df.iloc[0]["surplus_kwh"] == 120.0
    assert df.iloc[0]["grid_demand_kwh"] == 20.0

    assert df.iloc[1]["use"] == "Garaje"
    assert df.iloc[1]["share_pct"] == "40,0%"
    assert df.iloc[1]["consumption_kwh"] == 300.0
    assert df.iloc[1]["generation_kwh"] == 1000.0
    assert df.iloc[1]["self_consumed_kwh"] == 270.0
    assert df.iloc[1]["surplus_kwh"] == 130.0
    assert df.iloc[1]["grid_demand_kwh"] == 30.0


def test_export_all_filtering(tmp_path: Path) -> None:
    opt_result, meta = create_dummy_optimization_result()

    # CSV exports monthly_shares and hourly_shares
    paths_csv = ReportExporter(output_dir=tmp_path).export_all(opt_result, meta, formats="csv")
    assert len(paths_csv) == 2
    assert all(p.suffix == ".csv" for p in paths_csv)
    file_names = [p.name for p in paths_csv]
    assert any("_esd_monthly_shares.csv" in n for n in file_names)
    assert any("_esd_hourly_shares.csv" in n for n in file_names)

    # HTML exports interactive_report.html
    paths_html = ReportExporter(output_dir=tmp_path).export_all(opt_result, meta, formats="html")
    assert len(paths_html) == 1
    assert paths_html[0].name.endswith("_esd_interactive_report.html")

    # All exports monthly_shares, hourly_shares, and interactive_report
    paths_all = ReportExporter(output_dir=tmp_path).export_all(opt_result, meta, formats="all")
    assert len(paths_all) == 3

    # None / table
    paths_none = ReportExporter(output_dir=tmp_path).export_all(opt_result, meta, formats="table")
    assert len(paths_none) == 0


def test_export_hourly_files(tmp_path: Path) -> None:
    opt_result, meta = create_dummy_optimization_result()
    cups_list = ["ES0021000000000001AA", "ES0021000000000002BB"]
    hours = [f"{h:02d}" for h in range(24)]

    profile_matrix = HourlyProfileMatrix(
        cups_list=cups_list,
        hours=hours,
        matrix={
            cups_list[0]: dict.fromkeys(hours, 0.6),
            cups_list[1]: dict.fromkeys(hours, 0.4),
        },
        hourly_sums=dict.fromkeys(hours, 1.0),
        average_shares={cups_list[0]: 0.6, cups_list[1]: 0.4},
        average_sum=1.0,
    )
    ts_rows = []
    for h in hours:
        ts_rows.append(
            {
                "timestamp": f"2026-05-01 {h}:00:00",
                "hour": int(h),
                "cups": cups_list[0],
                "use": "Portal 1",
                "consumption_kwh": 2.4,
                "generation_kwh": 5.0,
                "share_pct": "60%",
                "estimated_share_pct": "60%",
                "solar_efficiency": 75,
            }
        )
        ts_rows.append(
            {
                "timestamp": f"2026-05-01 {h}:00:00",
                "hour": int(h),
                "cups": cups_list[1],
                "use": "Aerotermia",
                "consumption_kwh": 1.6,
                "generation_kwh": 5.0,
                "share_pct": "40%",
                "estimated_share_pct": "40%",
                "solar_efficiency": 75,
            }
        )
    ts_df = pd.DataFrame(ts_rows)
    opt_result.hourly_result = HourlyOptimizationResult(
        profile_matrix=profile_matrix,
        time_series_df=ts_df,
    )

    # Export hourly CSV file directly
    shares_path = ReportExporter(
        output_dir=tmp_path, timestamp="20260531_120000"
    ).export_hourly_shares_csv(opt_result)
    assert shares_path.exists()
    shares_content = shares_path.read_text(encoding="utf-8")
    assert (
        "timestamp;hour;cups;use;consumption_kwh;generation_kwh;share_pct;estimated_share_pct;solar_efficiency"
        in shares_content
    )
    assert "ES0021000000000001AA" in shares_content

    # Verify that within the same timestamp, Aerotermia appears before Portal 1
    lines = shares_content.strip().splitlines()
    assert len(lines) == 1 + (len(hours) * 2)
    assert "Aerotermia" in lines[1]
    assert "Portal 1" in lines[2]

    # export_all with formats='csv' should export 2 CSVs (monthly_shares and hourly_shares)
    exported = ReportExporter(output_dir=tmp_path).export_all(opt_result, meta, formats="csv")
    assert len(exported) == 2
    file_names = [p.name for p in exported]
    assert any("_esd_monthly_shares.csv" in n for n in file_names)
    assert any("_esd_hourly_shares.csv" in n for n in file_names)


def test_export_csvs_semicolon_separator_comma_decimal(tmp_path: Path) -> None:
    """Verify that all CSV outputs use ';' as separator and ',' for decimals without thousand separators."""
    opt_result, meta = create_dummy_optimization_result()
    cups_list = ["ES0021000000000001AA", "ES0021000000000002BB"]
    hours = [f"{h:02d}" for h in range(24)]

    profile_matrix = HourlyProfileMatrix(
        cups_list=cups_list,
        hours=hours,
        matrix={
            cups_list[0]: dict.fromkeys(hours, 0.625),
            cups_list[1]: dict.fromkeys(hours, 0.375),
        },
        hourly_sums=dict.fromkeys(hours, 1.0),
        average_shares={cups_list[0]: 0.625, cups_list[1]: 0.375},
        average_sum=1.0,
    )
    ts_rows = []
    for h in hours:
        ts_rows.append(
            {
                "timestamp": f"2026-05-01 {h}:00:00",
                "hour": int(h),
                "cups": cups_list[0],
                "use": "Heating",
                "consumption_kwh": 1234.56,
                "generation_kwh": 5000.78,
                "share_pct": "62.5%",
                "estimated_share_pct": "62.5%",
                "solar_efficiency": 82,
            }
        )
    opt_result.hourly_result = HourlyOptimizationResult(
        profile_matrix=profile_matrix,
        time_series_df=pd.DataFrame(ts_rows),
    )

    csv_paths = ReportExporter(output_dir=tmp_path).export_all(
        opt_result, meta, formats="csv", precision=2
    )
    assert len(csv_paths) == 2

    for p in csv_paths:
        text = p.read_text(encoding="utf-8")
        lines = text.strip().splitlines()
        assert ";" in lines[0]
        for line in lines[1:]:
            assert ";" in line
            for part in line.split(";"):
                assert not (part.count(",") > 1 or ("," in part and "." in part))

    # Check monthly shares CSV specific values
    monthly_path = next(p for p in csv_paths if "_esd_monthly_shares.csv" in p.name)
    monthly_text = monthly_path.read_text(encoding="utf-8")
    assert "60,00%" in monthly_text
    assert "40,00%" in monthly_text

    # Check hourly shares CSV specific values
    shares_path = next(p for p in csv_paths if "_esd_hourly_shares.csv" in p.name)
    shares_text = shares_path.read_text(encoding="utf-8")
    assert "1234,56" in shares_text
    assert "5000,78" in shares_text
    assert "62,5%" in shares_text
    assert "82" in shares_text


def test_report_exporter_class(tmp_path: Path) -> None:
    opt_result, meta = create_dummy_optimization_result()
    out_dir = tmp_path / "custom_output"
    exporter = ReportExporter(
        output_dir=out_dir,
        delimiter=";",
        decimal_char=",",
        timestamp="20261025_120000",
    )
    assert exporter.output_dir == out_dir
    assert exporter.delimiter == ";"
    assert exporter.decimal_char == ","
    assert exporter.timestamp == "20261025_120000"

    # Test all exports via ReportExporter instance
    exported = exporter.export_all(opt_result, meta, formats="all", precision=1)
    assert len(exported) == 3
    for p in exported:
        assert p.exists()
        assert p.parent == out_dir
        assert p.name.startswith("20261025_120000")


def test_month_abbr_consolidation_and_formatter() -> None:
    """Ensure MONTH_ABBR_EN/ES contain 12 months ('01'..'12') and format_month_abbr resolves cleanly."""
    expected_en = [
        "Jan",
        "Feb",
        "Mar",
        "Apr",
        "May",
        "Jun",
        "Jul",
        "Aug",
        "Sep",
        "Oct",
        "Nov",
        "Dec",
    ]
    expected_es = [
        "Ene",
        "Feb",
        "Mar",
        "Abr",
        "May",
        "Jun",
        "Jul",
        "Ago",
        "Sep",
        "Oct",
        "Nov",
        "Dic",
    ]

    assert len(MONTH_ABBR_EN) == 12
    assert len(MONTH_ABBR_ES) == 12

    for i in range(1, 13):
        k = f"{i:02d}"
        assert MONTH_ABBR_EN[k] == expected_en[i - 1]
        assert MONTH_ABBR_ES[k] == expected_es[i - 1]
        assert format_month_abbr(k, lang="en") == expected_en[i - 1]
        assert format_month_abbr(k, lang="es") == expected_es[i - 1]
        assert format_month_abbr(i, lang="en") == expected_en[i - 1]
        assert format_month_abbr(i, lang="es") == expected_es[i - 1]
        assert format_month_abbr(f"2026-{k}", lang="en") == expected_en[i - 1]
        assert format_month_abbr(f"2026-{k}", lang="es") == expected_es[i - 1]


def test_sanitize_formula_injection_helper() -> None:
    """sanitize_formula_injection escapes dangerous spreadsheet prefixes while preserving safe values."""
    assert sanitize_formula_injection("=SUM(1,2)") == "'=SUM(1,2)"
    assert sanitize_formula_injection("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_formula_injection("-MinusUse") == "'-MinusUse"
    assert sanitize_formula_injection("@SUM(A1:A10)") == "'@SUM(A1:A10)"
    assert sanitize_formula_injection("\tTabbed") == "'\tTabbed"
    assert sanitize_formula_injection("\rCarriage") == "'\rCarriage"

    assert sanitize_formula_injection("Aerotermia") == "Aerotermia"
    assert sanitize_formula_injection("Garaje") == "Garaje"
    assert sanitize_formula_injection("Community Lighting") == "Community Lighting"
    assert sanitize_formula_injection("Ascensor 1") == "Ascensor 1"
    assert sanitize_formula_injection("") == ""

    assert sanitize_formula_injection(None) is None
    assert sanitize_formula_injection(123) == 123
    assert sanitize_formula_injection(45.67) == 45.67
    assert sanitize_formula_injection(True) is True

    assert set(FORMULA_INJECTION_PREFIXES) == {"=", "+", "-", "@", "\t", "\r"}


def test_formula_injection_neutralized_in_monthly_and_hourly_shares_csv(tmp_path: Path) -> None:
    """Monthly shares and hourly shares CSV exports prefix formula-starting use labels with single quote."""
    c1 = CupsMonthlyMetrics(
        cups="ES0021000000000001AA",
        use="=2+2",
        month="2026-05",
        beta=0.6,
        consumption_kwh=500.0,
        generation_allocated_kwh=600.0,
        self_consumed_kwh=480.0,
        surplus_kwh=120.0,
        grid_demand_kwh=20.0,
        self_consumption_rate=80.0,
        solar_coverage_rate=96.0,
    )
    c2 = CupsMonthlyMetrics(
        cups="ES0021000000000002BB",
        use="@SUM(B1:B10)",
        month="2026-05",
        beta=0.4,
        consumption_kwh=300.0,
        generation_allocated_kwh=400.0,
        self_consumed_kwh=270.0,
        surplus_kwh=130.0,
        grid_demand_kwh=30.0,
        self_consumption_rate=67.5,
        solar_coverage_rate=90.0,
    )
    m1 = CommunityMonthlyMetrics(
        month="2026-05",
        strategy=OptimizationStrategy.OPTIMAL,
        total_generation_kwh=1000.0,
        total_consumption_kwh=800.0,
        total_self_consumed_kwh=750.0,
        total_surplus_kwh=250.0,
        total_grid_demand_kwh=50.0,
        self_consumption_rate=75.0,
        solar_coverage_rate=93.75,
        beta_sum=1.0,
        cups_metrics=[c1, c2],
    )

    time_series_df = pd.DataFrame(
        [
            {
                "timestamp": "2026-05-01 12:00:00",
                "cups": "ES0021000000000001AA",
                "use": "+cmd|' /C calc'!A0",
                "share_pct": "60.00",
                "estimated_share_pct": "60.00",
            },
            {
                "timestamp": "2026-05-01 12:00:00",
                "cups": "ES0021000000000002BB",
                "use": "Normal Use",
                "share_pct": "40.00",
                "estimated_share_pct": "40.00",
            },
        ]
    )

    profile_matrix = HourlyProfileMatrix(
        hours=[f"{h:02d}" for h in range(24)],
        cups_list=["ES0021000000000001AA", "ES0021000000000002BB"],
        matrix={"ES0021000000000001AA": {"12": 0.6}, "ES0021000000000002BB": {"12": 0.4}},
        hourly_sums={"12": 1.0},
        average_shares={"ES0021000000000001AA": 0.6, "ES0021000000000002BB": 0.4},
        average_sum=1.0,
        cups_uses={"ES0021000000000001AA": "-MinusFormula", "ES0021000000000002BB": "Normal Use"},
    )

    result = OptimizationResult(
        strategy=OptimizationStrategy.OPTIMAL,
        monthly_results=[m1],
        total_summary=m1,
        baselines={},
        hourly_result=HourlyOptimizationResult(
            time_series_df=time_series_df,
            profile_matrix=profile_matrix,
        ),
        cups_uses={"ES0021000000000001AA": "=2+2", "ES0021000000000002BB": "@SUM(B1:B10)"},
    )

    exporter = ReportExporter(output_dir=tmp_path, timestamp="20260531_test")

    # 1. Test export_monthly_shares_csv
    monthly_csv = exporter.export_monthly_shares_csv(result)
    raw_text = monthly_csv.read_text(encoding="utf-8")
    assert "'=2+2" in raw_text
    assert "'@SUM(B1:B10)" in raw_text

    df_monthly = pd.read_csv(monthly_csv, sep=";", decimal=",")
    assert set(df_monthly["use"].tolist()) == {"'=2+2", "'@SUM(B1:B10)"}

    # 2. Test export_hourly_shares_csv
    shares_csv = exporter.export_hourly_shares_csv(result)
    df_shares = pd.read_csv(shares_csv, sep=";")
    assert (
        df_shares[df_shares["cups"] == "ES0021000000000001AA"]["use"].iloc[0]
        == "'+cmd|' /C calc'!A0"
    )
    assert df_shares[df_shares["cups"] == "ES0021000000000002BB"]["use"].iloc[0] == "Normal Use"


def test_formula_injection_neutralized_in_curated_csv_and_excel(tmp_path: Path) -> None:
    """Curated CSV and Excel exports escape formula injection in the use column."""
    import openpyxl

    from src.curation.models import (
        CuratedCupsReading,
        CuratedDataset,
        CuratedHourRecord,
        CuratedSummary,
    )

    rec = CuratedHourRecord(
        timestamp="2026-05-01 12:00:00",
        date="2026-05-01",
        hour="12:00",
        generation=15.0,
        self_consumption=10.0,
        exported_energy=5.0,
        imported_energy=0.0,
        cups=[
            CuratedCupsReading(
                id="ES0021000000000001AA",
                use="=1+1",
                consumption=6.0,
            ),
            CuratedCupsReading(
                id="ES0021000000000002BB",
                use="Aerotermia",
                consumption=4.0,
            ),
        ],
    )

    dataset = CuratedDataset(
        records=[rec],
        summary=CuratedSummary(
            total_hours=1,
            reconciled_hours_count=0,
            recovered_kwh=0.0,
            total_generation_kwh=15.0,
            total_consumption_kwh=10.0,
            start_time="2026-05-01 12:00:00",
            end_time="2026-05-01 12:00:00",
            cups_count=2,
            primary_cups="ES0021000000000001AA",
        ),
        cups_configs={},
    )

    exporter = ReportExporter(output_dir=tmp_path, timestamp="20260531_curated")

    # 1. Curated CSV export
    csv_file = exporter.export_curated_dataset(dataset)
    raw_csv = csv_file.read_text(encoding="utf-8")
    assert "'=1+1" in raw_csv
    assert "Aerotermia" in raw_csv

    df_csv = pd.read_csv(csv_file, sep=";", decimal=",")
    assert df_csv[df_csv["cups_id"] == "ES0021000000000001AA"]["use"].iloc[0] == "'=1+1"
    assert df_csv[df_csv["cups_id"] == "ES0021000000000002BB"]["use"].iloc[0] == "Aerotermia"

    # 2. Curated Excel export
    excel_file = exporter.export_curated_excel(dataset)
    assert excel_file.exists()

    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active
    assert ws is not None

    headers = [cell.value for cell in ws[1]]
    use_col_idx = headers.index("use") + 1

    val_row2 = ws.cell(row=2, column=use_col_idx).value
    assert val_row2 == "'=1+1"

    val_row3 = ws.cell(row=3, column=use_col_idx).value
    assert val_row3 == "Aerotermia"


def test_export_interactive_html_includes_total_footer(tmp_path: Path) -> None:
    """Ensure ReportExporter.export_interactive_html produces matrix table with <tfoot> and 100.00%."""
    opt_result, meta = create_dummy_optimization_result()
    html_path = ReportExporter(output_dir=tmp_path).export_interactive_html(
        opt_result, meta, precision=2
    )
    assert html_path.exists()
    content = html_path.read_text(encoding="utf-8")
    assert "<tfoot>" in content
    assert "TOTAL (RD 244/2019)" in content
    assert "100.00%" in content
