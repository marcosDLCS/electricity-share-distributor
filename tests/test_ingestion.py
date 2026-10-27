"""Automated tests for DATADIS consumption and Huawei generation data ingestion."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.ingestion.consumption import DatadisConsumptionLoader
from src.ingestion.generation import HuaweiGenerationLoader
from src.ingestion.schema import IngestionError, ValidationError
from src.ingestion.weather import (
    ALLOWED_WEATHER_HOST,
    MAX_RESPONSE_BYTES,
    OpenMeteoWeatherDownloader,
)


def create_mock_datadis_csv(
    file_path: Path,
    cups: str = "ES0021000000000001AA",
    delimiter: str = ";",
    decimal_sep: str = ",",
    encoding: str = "utf-8",
    num_days: int = 2,
    include_spring_dst: bool = False,
    include_autumn_dst: bool = False,
) -> None:
    """Helper to generate synthetic mock DATADIS consumption CSV files."""
    lines = [f"cups{delimiter}fecha{delimiter}hora{delimiter}consumo_kWh{delimiter}metodoObtencion"]
    for d in range(1, num_days + 1):
        month = "10" if include_autumn_dst else "03"
        date_str = f"2026/{month}/{d:02d}"
        if include_spring_dst and d == 29:
            # 23 hours on spring DST day (02:00 skipped)
            hours = [f"{h:02d}:00" for h in range(1, 25) if h != 2]
        elif include_autumn_dst and d == 25:
            # 25 hours on autumn DST day (25 sequential intervals with repeated 02:00)
            hours = [f"{h:02d}:00" for h in range(1, 26)]
        else:
            hours = [f"{h:02d}:00" for h in range(1, 25)]

        for h in hours:
            val_str = f"0{decimal_sep}150"
            lines.append(
                f'"{cups}"{delimiter}"{date_str}"{delimiter}"{h}"{delimiter}"{val_str}"{delimiter}Real'
            )

    file_path.write_text("\n".join(lines), encoding=encoding)


def create_mock_huawei_excel(
    file_path: Path,
    num_days: int = 2,
    include_dst: bool = False,
) -> None:
    """Helper to generate synthetic mock Huawei FusionSolar Excel reports."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Row 1: Title
    ws.append(["Informe de plantas_MOCK INSTALLATION", "", "", "", "", ""])

    # Row 2: Headers
    ws.append(
        [
            "Período estadístico",
            "Irradiancia global (kWh/m²)",
            "Temperatura media (°C)",
            "Rendimiento teórico (kWh)",
            "Rendimiento FV (kWh)",
            "Rendimiento del inversor (kWh)",
        ]
    )

    # Data rows
    for d in range(1, num_days + 1):
        for h in range(24):
            dst_suffix = " DST" if include_dst else ""
            ts = f"2026-03-{d:02d} {h:02d}:00:00{dst_suffix}"
            gen_val = 5.5 if 9 <= h <= 17 else 0.0
            ws.append([ts, 0.5, 18.0, 6.0, gen_val, gen_val])

    wb.save(file_path)


def test_consumption_loader_directory_not_found(tmp_path: Path) -> None:
    loader = DatadisConsumptionLoader(tmp_path / "non_existent")
    with pytest.raises(IngestionError, match="not found"):
        loader.discover_files()


def test_consumption_loader_empty_directory(tmp_path: Path) -> None:
    loader = DatadisConsumptionLoader(tmp_path)
    with pytest.raises(IngestionError, match="No CSV files found"):
        loader.discover_files()


def test_consumption_loader_valid_files(tmp_path: Path) -> None:
    csv_file = tmp_path / "sample.csv"
    create_mock_datadis_csv(csv_file, delimiter=";", decimal_sep=",")

    loader = DatadisConsumptionLoader(tmp_path)
    df = loader.load_all()

    assert not df.empty
    assert len(df) == 48
    assert "cups" in df.columns
    assert "timestamp" in df.columns
    assert "consumption_kwh" in df.columns
    assert df["cups"].iloc[0] == "ES0021000000000001AA"
    assert df["consumption_kwh"].iloc[0] == 0.15


def test_consumption_loader_different_delimiters_and_encodings(tmp_path: Path) -> None:
    csv1 = tmp_path / "comma_dot.csv"
    create_mock_datadis_csv(
        csv1,
        cups="ES0021000000000001AA",
        delimiter=",",
        decimal_sep=".",
        encoding="latin-1",
    )

    loader = DatadisConsumptionLoader(tmp_path)
    df = loader.load_all()
    assert len(df) == 48
    assert df["consumption_kwh"].iloc[0] == 0.15


def test_consumption_loader_missing_required_column(tmp_path: Path) -> None:
    bad_csv = tmp_path / "bad.csv"
    bad_csv.write_text("fecha;hora;consumo_kWh\n2026/03/01;01:00;0.1", encoding="utf-8")

    loader = DatadisConsumptionLoader(tmp_path)
    with pytest.raises(ValidationError, match="missing required DATADIS column: 'cups'"):
        loader.parse_file(bad_csv)


def test_consumption_spring_dst_transition(tmp_path: Path) -> None:
    csv_file = tmp_path / "dst.csv"
    create_mock_datadis_csv(csv_file, num_days=29, include_spring_dst=True)

    loader = DatadisConsumptionLoader(tmp_path)
    df = loader.parse_file(csv_file)

    # March 29 should have 23 readings
    m29 = df[df["timestamp"].str.startswith("2026-03-29")]
    assert len(m29) == 23
    # Check that timestamps map 00:00:00, 01:00:00, 03:00:00, etc.
    ts_list = m29["timestamp"].tolist()
    assert ts_list[0] == "2026-03-29 00:00:00"
    assert ts_list[1] == "2026-03-29 01:00:00"
    assert ts_list[2] == "2026-03-29 03:00:00"


def test_consumption_autumn_dst_transition(tmp_path: Path) -> None:
    csv_file = tmp_path / "dst_autumn.csv"
    create_mock_datadis_csv(csv_file, num_days=25, include_autumn_dst=True)

    loader = DatadisConsumptionLoader(tmp_path)
    df = loader.parse_file(csv_file)

    # October 25 should have 25 readings
    oct25 = df[df["timestamp"].str.startswith("2026-10-25")]
    assert len(oct25) == 25
    ts_list = oct25["timestamp"].tolist()
    assert ts_list[0] == "2026-10-25 00:00:00"
    assert ts_list[1] == "2026-10-25 01:00:00"
    assert ts_list[2] == "2026-10-25 02:00:00"
    assert ts_list[3] == "2026-10-25 02:00:00"
    assert ts_list[4] == "2026-10-25 03:00:00"
    assert ts_list[24] == "2026-10-25 23:00:00"


def test_consumption_unparseable_hora_fallback(tmp_path: Path) -> None:
    csv_file = tmp_path / "unparseable_hora.csv"
    # Create CSV where hora has malformed strings like 'invalid' or '??'
    lines = [
        "cups;fecha;hora;consumo_kWh;metodoObtencion",
        "ES0021000000000001AA;2026/04/01;invalid;0,150;Real",
        "ES0021000000000001AA;2026/04/01;02:00;0,150;Real",
        "ES0021000000000001AA;2026/04/01;??;0,150;Real",
    ]
    csv_file.write_text("\n".join(lines), encoding="utf-8")

    loader = DatadisConsumptionLoader(tmp_path)
    df = loader.parse_file(csv_file)
    assert len(df) == 3
    # idx 0 fallback: h = idx + 1 = 1 -> 00:00:00
    # idx 1: h = 2 -> 01:00:00
    # idx 2 fallback: h = idx + 1 = 3 -> 02:00:00
    assert df["timestamp"].tolist() == [
        "2026-04-01 00:00:00",
        "2026-04-01 01:00:00",
        "2026-04-01 02:00:00",
    ]


def test_generation_loader_valid(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "Informe_planta_test.xlsx"
    create_mock_huawei_excel(xlsx_file, num_days=2, include_dst=True)

    loader = HuaweiGenerationLoader(tmp_path)
    df = loader.load_all()

    assert not df.empty
    assert len(df) == 48
    assert "timestamp" in df.columns
    assert "generation_kwh" in df.columns
    # Ensure DST suffix is removed
    assert not df["timestamp"].str.contains("DST").any()
    assert df["generation_kwh"].max() == 5.5


def test_generation_loader_missing_columns(tmp_path: Path) -> None:
    bad_xlsx = tmp_path / "bad_gen.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Title"])
    ws.append(["Fecha", "Otra columna"])
    ws.append(["2026-01-01 00:00:00", 123])
    wb.save(bad_xlsx)

    loader = HuaweiGenerationLoader(tmp_path)
    with pytest.raises(ValidationError, match="does not contain 'Rendimiento FV'"):
        loader.parse_file(bad_xlsx)


def test_generation_loader_with_autoconsumo(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "Informe_planta_autoconsumo.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Informe de plantas_TEST", "", "", "", "", "", ""])
    ws.append(
        [
            "Período estadístico",
            "Rendimiento FV (kWh)",
            "Consumo (kWh)",
            "Autoconsumo (kWh)",
            "Energía exportada (kWh)",
            "Importación (kWh)",
        ]
    )
    ws.append(["2026-06-01 12:00:00", 10.5, 8.0, 7.5, 3.0, 0.5])
    wb.save(xlsx_file)

    loader = HuaweiGenerationLoader(tmp_path)
    df = loader.load_all()

    assert not df.empty
    assert "generation_kwh" in df.columns
    assert "consumption_kwh" in df.columns
    assert "self_consumption_kwh" in df.columns
    assert "exported_energy_kwh" in df.columns
    assert "imported_energy_kwh" in df.columns
    assert df["self_consumption_kwh"].iloc[0] == 7.5
    assert df["exported_energy_kwh"].iloc[0] == 3.0
    assert df["imported_energy_kwh"].iloc[0] == 0.5


def test_weather_downloader_security_host_and_payload_limits(tmp_path: Path) -> None:
    """OpenMeteoWeatherDownloader enforces HTTPS, trusted hostname, and 10 MB payload limits."""
    from unittest.mock import MagicMock, patch

    # Insecure scheme rejected
    with pytest.raises(IngestionError, match="Only HTTPS is permitted"):
        OpenMeteoWeatherDownloader.validate_url("http://archive-api.open-meteo.com/v1/archive")

    # SSRF untrusted host rejected
    with pytest.raises(IngestionError, match=r"Untrusted host 'attacker\.com'"):
        OpenMeteoWeatherDownloader(weather_dir=tmp_path, base_url="https://attacker.com/v1/archive")

    # Official endpoint accepted
    OpenMeteoWeatherDownloader.validate_url(
        f"https://{ALLOWED_WEATHER_HOST}/v1/archive?latitude=40.4&longitude=-3.7"
    )

    # Oversized payload rejected (> 10 MB)
    downloader = OpenMeteoWeatherDownloader(weather_dir=tmp_path)
    mock_resp = MagicMock()
    mock_resp.read.return_value = b"A" * (MAX_RESPONSE_BYTES + 1)
    mock_resp.__enter__.return_value = mock_resp
    mock_resp.__exit__.return_value = None

    with patch("urllib.request.urlopen", return_value=mock_resp):
        with pytest.raises(IngestionError, match="exceeded maximum allowable size"):
            downloader.download_year(year=2024, latitude=40.4168, longitude=-3.7038)


def test_consumption_loader_spanish_date_formats(tmp_path: Path) -> None:
    """Verify that Spanish DD/MM/YYYY and DD-MM-YYYY dates parse into identical timestamps as YYYY/MM/DD and YYYY-MM-DD."""
    cups = "ES0021000000000001AA"
    hours = ["01:00", "02:00", "12:00", "24:00"]
    expected_ts = [
        "2026-03-15 00:00:00",
        "2026-03-15 01:00:00",
        "2026-03-15 11:00:00",
        "2026-03-15 23:00:00",
    ]

    formats = {
        "slash_dmy": "15/03/2026",
        "dash_dmy": "15-03-2026",
        "slash_ymd": "2026/03/15",
        "dash_ymd": "2026-03-15",
    }

    loader = DatadisConsumptionLoader(tmp_path)

    for fmt_name, date_str in formats.items():
        file_p = tmp_path / f"test_{fmt_name}.csv"
        lines = ["cups;fecha;hora;consumo_kWh;metodoObtencion"]
        for h in hours:
            lines.append(f'"{cups}";"{date_str}";"{h}";"0,250";Real')
        file_p.write_text("\n".join(lines), encoding="utf-8")

        parsed_df = loader.parse_file(file_p)
        assert len(parsed_df) == 4
        assert parsed_df["timestamp"].tolist() == expected_ts
        assert parsed_df["cups"].tolist() == [cups] * 4
        assert parsed_df["consumption_kwh"].tolist() == [0.25] * 4


def test_consumption_loader_single_digit_spanish_dates(tmp_path: Path) -> None:
    """Verify that single-digit day and month in Spanish formats are normalized with leading zeros."""
    cups = "ES0021000000000001AA"
    loader = DatadisConsumptionLoader(tmp_path)

    cases = [
        ("5/4/2026", "2026-04-05 00:00:00"),
        ("5-4-2026", "2026-04-05 00:00:00"),
        ("2026-4-5", "2026-04-05 00:00:00"),
        ("2026/4/5", "2026-04-05 00:00:00"),
    ]

    for i, (date_str, expected_ts) in enumerate(cases):
        file_p = tmp_path / f"single_{i}.csv"
        lines = [
            "cups;fecha;hora;consumo_kWh;metodoObtencion",
            f'"{cups}";"{date_str}";"01:00";"0,500";Real',
        ]
        file_p.write_text("\n".join(lines), encoding="utf-8")
        parsed = loader.parse_file(file_p)
        assert parsed["timestamp"].iloc[0] == expected_ts
