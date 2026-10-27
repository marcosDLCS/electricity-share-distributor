"""Unit tests for the security and privacy validation harness."""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

import openpyxl

from src.security import (
    check_privacy,
    clear_private_signatures_cache,
    find_cups_leaks,
    find_path_leaks,
    get_known_private_signatures,
    is_synthetic_cups,
    scan_directory,
    scan_file,
    scan_git_history,
    scan_git_staged,
    scan_xlsx_file,
)

# Test fixture string constructed dynamically so it does not contain any real CUPS
PROHIBITED_SAMPLE_CUPS = "ES" + "9" * 16 + "ZZ"
PROHIBITED_SAMPLE_CUPS_2 = "ES" + "8" * 16 + "WW"


def test_is_synthetic_cups_valid() -> None:
    assert is_synthetic_cups("ES0021000000000001AA")
    assert is_synthetic_cups("ES0021000000000002BB")
    assert is_synthetic_cups("ES0021000000000009II")
    assert is_synthetic_cups("es0021000000000001aa")


def test_is_synthetic_cups_real_detected() -> None:
    # Non-synthetic CUPS format must return False
    assert not is_synthetic_cups(PROHIBITED_SAMPLE_CUPS)
    assert not is_synthetic_cups("ES" + "1234567890123456" + "AB")


def test_find_cups_leaks_clean() -> None:
    clean_text = "Participant 1: ES0021000000000001AA, Participant 2: ES0021000000000002BB"
    assert find_cups_leaks(clean_text) == []


def test_find_cups_leaks_dirty() -> None:
    dirty_text = f"Prohibited CUPS in text: {PROHIBITED_SAMPLE_CUPS} and ES0021000000000001AA"
    leaks = find_cups_leaks(dirty_text)
    assert leaks == [PROHIBITED_SAMPLE_CUPS]


def test_find_path_leaks() -> None:
    clean_text = "Relative link: [file](../src/cli.py) or `src/cli.py`"
    assert find_path_leaks(clean_text) == []

    # Prohibited local file paths
    fake_path = "/Users" + "/john/workspace/project"
    dirty_text = f"Local file: {fake_path} and file:///Users/someone/doc.md"
    leaks = find_path_leaks(dirty_text)
    assert len(leaks) >= 1


def test_scan_file(tmp_path: Path) -> None:
    file_clean = tmp_path / "clean.md"
    file_clean.write_text("All good with ES0021000000000001AA\n")
    assert scan_file(file_clean) == []

    file_dirty = tmp_path / "dirty.md"
    file_dirty.write_text(f"Header\nLeak: {PROHIBITED_SAMPLE_CUPS}\nFooter\n")
    findings = scan_file(file_dirty)
    assert len(findings) == 1
    assert findings[0][0] == 2  # Line 2
    assert findings[0][1] == PROHIBITED_SAMPLE_CUPS


def test_get_known_private_signatures_from_cups_json(tmp_path: Path) -> None:
    """Verify that get_known_private_signatures extracts non-synthetic CUPS from cups.json."""
    cups_dir = tmp_path / ".input" / "cups"
    cups_dir.mkdir(parents=True, exist_ok=True)
    cups_file = cups_dir / "cups.json"

    data = [
        {"cups": "ES0021000000000001AA", "use": "piscina", "primary": False},
        {"cups": PROHIBITED_SAMPLE_CUPS_2, "use": "aerotermia", "primary": True},
    ]
    cups_file.write_text(json.dumps(data), encoding="utf-8")

    sigs = get_known_private_signatures(tmp_path)
    assert PROHIBITED_SAMPLE_CUPS_2 in sigs
    assert PROHIBITED_SAMPLE_CUPS_2[8:] in sigs
    assert "ES0021000000000001AA" not in sigs


def test_get_known_private_signatures_from_flat_cups_json(tmp_path: Path) -> None:
    """Verify discovery from fallback .input/cups.json location."""
    input_dir = tmp_path / ".input"
    input_dir.mkdir(parents=True, exist_ok=True)
    cups_file = input_dir / "cups.json"

    data = {"primary_point": PROHIBITED_SAMPLE_CUPS_2}
    cups_file.write_text(json.dumps(data), encoding="utf-8")

    sigs = get_known_private_signatures(tmp_path)
    assert PROHIBITED_SAMPLE_CUPS_2 in sigs
    assert PROHIBITED_SAMPLE_CUPS_2[8:] in sigs


def test_clear_private_signatures_cache() -> None:
    """Verify cache clearing function resets the global cache."""
    clear_private_signatures_cache()
    # Cache is cleared, next call recalculates
    sigs = get_known_private_signatures()
    assert isinstance(sigs, set)


def test_scan_xlsx_file_clean(tmp_path: Path) -> None:
    """Verify scan_xlsx_file returns no findings for clean workbook with synthetic CUPS."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Header"
    ws["B1"] = "ES0021000000000001AA"
    ws["C1"] = 45.67

    xlsx_file = tmp_path / "clean.xlsx"
    wb.save(xlsx_file)
    wb.close()

    assert scan_xlsx_file(xlsx_file) == []
    assert scan_file(xlsx_file) == []


def test_scan_xlsx_file_dirty(tmp_path: Path) -> None:
    """Verify scan_xlsx_file flags prohibited non-synthetic CUPS in cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PlantData"
    ws["B3"] = f"Facility {PROHIBITED_SAMPLE_CUPS} load"

    xlsx_file = tmp_path / "dirty.xlsx"
    wb.save(xlsx_file)
    wb.close()

    findings = scan_xlsx_file(xlsx_file)
    assert len(findings) == 1
    assert findings[0][0] == 3  # Row 3
    assert findings[0][1] == PROHIBITED_SAMPLE_CUPS
    assert "[PlantData] Row 3:" in findings[0][2]

    # scan_file must also delegate to scan_xlsx_file
    delegated = scan_file(xlsx_file)
    assert len(delegated) == 1
    assert delegated[0][1] == PROHIBITED_SAMPLE_CUPS


def test_scan_xlsx_file_path_leak(tmp_path: Path) -> None:
    """Verify scan_xlsx_file detects local filesystem paths in cell values."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Paths"
    fake_path = "/Users" + "/developer/code/esd"
    ws["A1"] = fake_path

    xlsx_file = tmp_path / "paths.xlsx"
    wb.save(xlsx_file)
    wb.close()

    findings = scan_xlsx_file(xlsx_file)
    assert len(findings) >= 1
    assert fake_path in findings[0][1]


def test_scan_xlsx_file_invalid_or_missing(tmp_path: Path) -> None:
    """Verify scan_xlsx_file gracefully returns empty list for non-existent or corrupt files."""
    missing = tmp_path / "non_existent.xlsx"
    assert scan_xlsx_file(missing) == []

    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_text("not a valid zip", encoding="utf-8")
    assert scan_xlsx_file(corrupt) == []


def test_scan_directory_detects_xlsx_outside_ignored_dirs(tmp_path: Path) -> None:
    """Verify scan_directory inspects .xlsx files outside ignored dirs and ignores .input/.curated."""
    # 1. Dirty .xlsx in a tracked directory
    tracked_dir = tmp_path / "docs"
    tracked_dir.mkdir(parents=True, exist_ok=True)
    tracked_wb = openpyxl.Workbook()
    tracked_wb.active["A1"] = PROHIBITED_SAMPLE_CUPS
    tracked_file = tracked_dir / "example.xlsx"
    tracked_wb.save(tracked_file)
    tracked_wb.close()

    # 2. Dirty .xlsx inside .input/ (ignored)
    input_dir = tmp_path / ".input" / "generation"
    input_dir.mkdir(parents=True, exist_ok=True)
    input_wb = openpyxl.Workbook()
    input_wb.active["A1"] = PROHIBITED_SAMPLE_CUPS
    input_file = input_dir / "gen.xlsx"
    input_wb.save(input_file)
    input_wb.close()

    # 3. Dirty .xlsx inside .curated/ (ignored)
    curated_dir = tmp_path / ".curated"
    curated_dir.mkdir(parents=True, exist_ok=True)
    curated_wb = openpyxl.Workbook()
    curated_wb.active["A1"] = PROHIBITED_SAMPLE_CUPS
    curated_file = curated_dir / "curated.xlsx"
    curated_wb.save(curated_file)
    curated_wb.close()

    findings = scan_directory(tmp_path)
    assert tracked_file in findings
    assert findings[tracked_file][0][1] == PROHIBITED_SAMPLE_CUPS
    assert input_file not in findings
    assert curated_file not in findings


def test_scan_git_staged_detects_staged_xlsx(tmp_path: Path) -> None:
    """Verify scan_git_staged inspects staged .xlsx files in the git index."""
    repo = tmp_path / "test_repo"
    repo.mkdir()
    subprocess.run(["git", "init"], cwd=repo, check=True, capture_output=True)
    subprocess.run(["git", "config", "user.name", "Tester"], cwd=repo, check=True)
    subprocess.run(["git", "config", "user.email", "test@example.com"], cwd=repo, check=True)

    # 1. Create and stage a dirty .xlsx
    dirty_wb = openpyxl.Workbook()
    dirty_wb.active["B2"] = PROHIBITED_SAMPLE_CUPS
    dirty_file = repo / "staged_dirty.xlsx"
    dirty_wb.save(dirty_file)
    dirty_wb.close()

    subprocess.run(["git", "add", "staged_dirty.xlsx"], cwd=repo, check=True)

    findings = scan_git_staged(repo)
    assert "staged_dirty.xlsx" in findings
    assert findings["staged_dirty.xlsx"][0][0] == 2  # Row 2
    assert findings["staged_dirty.xlsx"][0][1] == PROHIBITED_SAMPLE_CUPS

    # 2. Stage a file in .input/ (ignored)
    input_dir = repo / ".input"
    input_dir.mkdir()
    input_wb = openpyxl.Workbook()
    input_wb.active["A1"] = PROHIBITED_SAMPLE_CUPS
    input_file = input_dir / "input_dirty.xlsx"
    input_wb.save(input_file)
    input_wb.close()

    subprocess.run(["git", "add", "-f", ".input/input_dirty.xlsx"], cwd=repo, check=True)
    findings2 = scan_git_staged(repo)
    assert ".input/input_dirty.xlsx" not in findings2


def test_entire_repository_is_clean() -> None:
    """Ensure the current workspace repository is 100% free of real CUPS identifiers."""
    repo_root = Path.cwd()
    findings = scan_directory(repo_root)
    assert findings == {}, f"Prohibited CUPS found in files: {findings}"


def test_entire_git_history_is_clean() -> None:
    """Ensure all commits in git history are completely free of real CUPS identifiers."""
    history_findings = scan_git_history(Path.cwd())
    assert history_findings == [], f"Prohibited CUPS found in git history: {history_findings}"


def test_check_privacy_succeeds() -> None:
    """Ensure the check_privacy entry point returns 0 for a clean repository."""
    assert check_privacy(include_history=True) == 0
    assert check_privacy(staged_only=True) == 0
